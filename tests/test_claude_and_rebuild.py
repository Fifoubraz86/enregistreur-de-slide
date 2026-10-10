import json
import sys
from pathlib import Path
from types import SimpleNamespace

import cv2
import numpy as np
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from capture_reunion import rebuild
from capture_reunion import transcribe as tr_mod
from capture_reunion.claude_code import ClaudeCode
from capture_reunion.llm import LLMError, parse_json
from capture_reunion.session import Session, SlideRecord
from capture_reunion.transcribe import transcribe_session
from capture_reunion.workflow import make_reports, rebuild_presentation
from slide_fixtures import chart_slide

FAKE = str(Path(__file__).with_name("fake_claude.py"))
pytestmark = pytest.mark.skipif(sys.platform == "win32", reason="faux exécutable POSIX")


@pytest.fixture
def claude_log(tmp_path, monkeypatch):
    log = tmp_path / "claude_calls.jsonl"
    monkeypatch.setenv("FAKE_CLAUDE_LOG", str(log))
    return log


def calls(log):
    return [json.loads(l) for l in log.read_text().splitlines()] if log.exists() else []


def test_parse_json_variants():
    assert parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json('Voici : {"a": [1, 2]} merci') == {"a": [1, 2]}
    assert parse_json("pas de json") is None


def test_claude_connect_and_flags(claude_log):
    c = ClaudeCode("sonnet", executable=FAKE)
    assert "Claude sonnet" in c.connect()
    assert c.chat("Consigne", "Bonjour") == "prêt"
    args = calls(claude_log)[0]["args"]
    assert args[:2] == ["-p", "Réalise la tâche décrite dans le texte transmis."]
    assert args[args.index("--model") + 1] == "sonnet"
    assert args[args.index("--output-format") + 1] == "json"
    assert "--system-prompt" in args  # réduit la consommation quand c'est accepté


def test_claude_retries_without_unsupported_flag(claude_log, monkeypatch):
    monkeypatch.setenv("FAKE_CLAUDE_REJECT_SYSTEM", "1")
    c = ClaudeCode("haiku", executable=FAKE)
    assert c.chat("Consigne", "Bonjour") == "prêt"
    assert "--system-prompt" not in calls(claude_log)[-1]["args"]
    assert c.chat("Consigne", "Encore") == "prêt"  # l'option n'est plus retentée


def test_claude_missing_executable(monkeypatch):
    monkeypatch.setattr("capture_reunion.claude_code.find_claude", lambda: None)
    with pytest.raises(LLMError, match="Claude Code est introuvable"):
        ClaudeCode().connect()


SCRIPT = [(0.0, 6.0, " Neuroendocrine tumours have a long course."),
          (6.0, 12.0, " We extend the lifespan of our patients."),
          (45.0, 51.0, " The bulk of the tumour is inside the liver.")]


@pytest.fixture
def session(tmp_path, monkeypatch):
    monkeypatch.setattr(tr_mod, "load_audio", lambda p: np.zeros(60 * 16000, np.float32))
    (tmp_path / "piste_participants.m4a").write_bytes(b"")
    (tmp_path / "diapos").mkdir()
    for i in (1, 2):
        ok, buf = cv2.imencode(".png", chart_slide())
        buf.tofile(str(tmp_path / "diapos" / f"diapo_00{i}.png"))
    s = Session(folder=tmp_path, started_at="2026-10-10T10:00:00", duration=60,
                tracks={"participants": "piste_participants.m4a"}, window_title="Webinar")
    s.slides = [SlideRecord(1, 0.0, "diapos/diapo_001.png"), SlideRecord(2, 40.0, "diapos/diapo_002.png")]
    s.timeline = [(0.0, 1), (40.0, 2)]
    s.save()
    return s


class Whisper:
    def transcribe(self, audio, **kwargs):
        return iter([SimpleNamespace(start=a, end=b, text=t) for a, b, t in SCRIPT]), \
            SimpleNamespace(language="en")


def test_transcription_with_claude_uses_few_calls(session, claude_log):
    transcribe_session(session, whisper=Whisper(), translate=True, engine="claude",
                       llm=ClaudeCode("sonnet", executable=FAKE))
    loaded = Session.load(session.folder)
    assert loaded.ai["model"].startswith("Claude sonnet")
    assert loaded.ai["slide_summaries"]["en"] == {"1": "Résumé de Slide 1.", "2": "Résumé de Slide 2."}
    assert loaded.ai["summary"]["fr"].startswith("## Sujet")
    fr_txt = (session.folder / "transcription_fr.txt").read_text(encoding="utf-8")
    assert "Traduction française de « Neuroendocrine tumours have »" in fr_txt
    # Résumés en 1 appel, traduction (paragraphes + phrases clés + « En bref ») en 1, résumé global en 1.
    assert len(calls(claude_log)) == 3
    assert [p.name for p in make_reports(loaded)] == ["compte-rendu.docx", "compte-rendu_fr.docx"]


def test_patient_data_blocks_claude(session, claude_log):
    transcribe_session(session, whisper=Whisper(), engine="claude", allow_online=False)
    loaded = Session.load(session.folder)
    assert calls(claude_log) == []
    assert "données patients" in loaded.ai["warnings"][0]


def test_ocr_reads_title_bullets_and_figure():
    c = rebuild.ocr_slide(chart_slide(), 1)
    assert c.title == "Neuroendocrine tumours: treatment"
    assert c.bullets == [("Somatostatin analogues first line", 0), ("Lanreotide or octreotide", 1),
                         ("PRRT after progression", 0), ("Liver-directed therapy", 0)]
    assert len(c.figures) == 1
    x, y, w, h = c.figures[0]
    assert 0.55 < x < 0.65 and w > 0.3 and y + h > 0.8  # le graphique, légende d'axe comprise


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def test_rebuild_with_ocr_and_default_template(session):
    out, warnings = rebuild_presentation(session)
    prs = Presentation(str(out))
    assert len(prs.slides) == 2 and warnings == []
    slide = prs.slides[0]
    texts = _texts(slide)
    assert "Neuroendocrine tumours: treatment" in texts
    body = next(sh for sh in slide.shapes if sh.has_text_frame and "Lanreotide" in sh.text_frame.text)
    levels = [(p.text, p.level) for p in body.text_frame.paragraphs]
    assert ("Lanreotide or octreotide", 1) in levels and ("PRRT after progression", 0) in levels
    pictures = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1 and pictures[0].left > prs.slide_width / 2
    pic = pictures[0]
    assert pic.width > prs.slide_width * 0.25 and pic.height > prs.slide_height * 0.3
    assert pic.top + pic.height <= prs.slide_height and pic.left + pic.width <= prs.slide_width
    title = next(sh for sh in slide.shapes if sh.has_text_frame and "treatment" in sh.text_frame.text)
    assert title.top > 0 and title.height > prs.slide_height * 0.08  # titre visible, pas écrasé
    assert body.width > prs.slide_width * 0.3 and body.height > prs.slide_height * 0.3
    assert (session.folder / "diapos_figures" / "diapo_001_figure_1.png").exists()
    assert "lecture : OCR local" in slide.notes_slide.notes_text_frame.text
    assert prs.slide_width > prs.slide_height  # masque par défaut 16:9
    numbers = [ph for ph in slide.placeholders if ph.placeholder_format.type is not None
               and "SLIDE_NUMBER" in str(ph.placeholder_format.type)]
    assert numbers and numbers[0].text_frame.text == "1"
    assert b'type="slidenum"' in numbers[0]._element.xml.encode()


def test_rebuild_with_claude_vision_and_cache(session, claude_log, monkeypatch):
    out, warnings = rebuild_presentation(session, llm=ClaudeCode("sonnet", executable=FAKE))
    assert warnings == []
    texts = _texts(Presentation(str(out)).slides[1])
    assert "Titre lu par Claude 2" in texts
    read_calls = [c for c in calls(claude_log) if "--allowedTools" in c["args"]]
    assert len(read_calls) == 1  # 2 diapos en un seul appel
    assert read_calls[0]["args"][read_calls[0]["args"].index("--allowedTools") + 1] == "Read"
    assert Path(read_calls[0]["cwd"]).resolve() == (session.folder / "diapos").resolve()
    # Deuxième fois (autre masque) : relu depuis le cache, aucun appel ni OCR.
    monkeypatch.setattr(rebuild, "_ocr", lambda image: (_ for _ in ()).throw(AssertionError("OCR")))
    n = len(calls(claude_log))
    rebuild_presentation(session, llm=ClaudeCode("sonnet", executable=FAKE))
    assert len(calls(claude_log)) == n


def test_rebuild_falls_back_to_ocr_when_ai_fails(session):
    class Broken:
        label = "Claude"
        display_name = "Claude"

        def connect(self):
            raise LLMError("non connecté")

    out, warnings = rebuild_presentation(session, llm=Broken())
    assert "OCR local utilisé" in warnings[0]
    assert "Neuroendocrine tumours: treatment" in _texts(Presentation(str(out)).slides[0])


def test_rebuild_with_user_template(session, tmp_path):
    from test_export import _make_template

    out, _ = rebuild_presentation(session, template=_make_template(tmp_path))
    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    assert "BANDEAU MAISON" in [sh.text_frame.text for sh in prs.slide_master.shapes if sh.has_text_frame]


def test_restore_accents():
    from capture_reunion.glossary import default_path, load

    vocab = rebuild.accent_vocabulary(load(default_path()))
    fixed = rebuild.restore_accents("Apres correction ; Arret des antagonistes mineralocorticoides. TRES", vocab)
    assert fixed == "Après correction ; Arrêt des antagonistes minéralocorticoïdes. TRÈS"
    assert rebuild.restore_accents("the liver after treatment", vocab) == "the liver after treatment"


def test_accents_only_on_french_slides():
    assert rebuild.looks_french("Rapport aldostérone / rénine après correction de l'hypokaliémie")
    assert not rebuild.looks_french("Somatostatin analogues first line. Lanreotide or octreotide")


def test_french_slide_reconstruction_fixes_ocr_glitches(tmp_path):
    from PIL import Image, ImageDraw

    from slide_fixtures import _font

    img = Image.new("RGB", (1280, 720), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 40), "Hyperaldostéronisme primaire : dépistage", fill=(20, 40, 90), font=_font(46))
    for k, (t, ind) in enumerate((("Rapport aldostérone / rénine", 0),
                                  ("Après correction de l'hypokaliémie", 1),
                                  ("Arrêt des antagonistes minéralocorticoïdes", 1))):
        d.text((80 + ind * 50, 170 + 62 * k), f"• {t}", fill="black", font=_font(30))
    c = rebuild.ocr_slide(cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR), 1)
    texts = [t for t, _ in c.bullets]
    assert "Après correction de l'hypokaliémie" in texts
    assert "Arrêt des antagonistes minéralocorticoïdes" in texts
