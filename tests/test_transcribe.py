"""Transcription locale, avec un faux modèle Whisper (le vrai se télécharge sur le PC)."""

import subprocess
from types import SimpleNamespace

import docx
import numpy as np
import pytest

from capture_reunion import transcribe as tr_mod
from capture_reunion.ffmpeg_utils import find_ffmpeg

from capture_reunion.session import Session, SlideRecord
from capture_reunion.transcribe import Track, remove_echo, transcribe_session, transcribe_tracks
from capture_reunion.transcript import Segment, load_transcript
from capture_reunion.workflow import make_report, make_reports

SCRIPTS = {
    "piste_participants.m4a": [(1.0, 4.0, " Bonjour à tous, on commence."),
                               (12.0, 15.0, " Voici les résultats.")],
    # Le micro capte sa propre voix, et l'écho des haut-parleurs à 12 s.
    "piste_moi.m4a": [(6.0, 9.0, " Merci, j'ai une question."),
                      (12.1, 15.0, " voici les résultats"),
                      (20.0, 22.0, " ")],
}


class FakeAudio(np.ndarray):
    """Tableau audio qui retient de quel fichier il vient."""


def fake_load_audio(path):
    audio = np.zeros(30 * tr_mod.SAMPLE_RATE, np.float32).view(FakeAudio)
    audio.name = path.name
    return audio


@pytest.fixture(autouse=True)
def no_ffmpeg_decode(monkeypatch, request):
    if "real_decode" not in request.keywords:
        monkeypatch.setattr(tr_mod, "load_audio", fake_load_audio)


class FakeWhisper:
    def __init__(self):
        self.calls = []

    def transcribe(self, audio, **kwargs):
        # Régression : on transmet le son déjà décodé, jamais un chemin (décodage PyAV).
        assert isinstance(audio, np.ndarray)
        self.calls.append((audio.name, kwargs))
        name = audio.name
        segs = [SimpleNamespace(start=a, end=b, text=t) for a, b, t in SCRIPTS[name]]
        return iter(segs), SimpleNamespace(duration=30.0, language="fr")


def test_remove_echo():
    others = [Segment(10, "Voici les résultats de l'étude", "Participants", 14)]
    mine = [Segment(10.2, "voici les résultats de l'étude", "Moi", 14),
            Segment(10.5, "Tout à fait d'accord", "Moi", 12)]
    assert [s.text for s in remove_echo(mine, others)] == ["Tout à fait d'accord"]


def test_transcribe_tracks_labels_and_progress(tmp_path):
    for name in SCRIPTS:
        (tmp_path / name).write_bytes(b"")
    progress = []
    whisper = FakeWhisper()
    segs, language = transcribe_tracks(
        [Track(tmp_path / "piste_participants.m4a", "Participants"), Track(tmp_path / "piste_moi.m4a", "Moi")],
        vocabulary="RCP, pembrolizumab", progress=progress.append, whisper=whisper)
    assert [(s.speaker, s.text) for s in segs] == [
        ("Participants", "Bonjour à tous, on commence."),
        ("Moi", "Merci, j'ai une question."),
        ("Participants", "Voici les résultats."),
    ]
    assert progress[-1] == 1.0 and progress == sorted(progress)
    kwargs = whisper.calls[0][1]
    # Régression : la langue n'est pas imposée (sinon une conférence en anglais est traduite).
    assert kwargs["language"] is None and kwargs["vad_filter"] is True
    assert language == "fr"
    assert kwargs["hotwords"] == "RCP, pembrolizumab"


def test_transcribe_session_feeds_report(tmp_path):
    for name in SCRIPTS:
        (tmp_path / name).write_bytes(b"")
    session = Session(folder=tmp_path, started_at="2026-10-03T14:30:00", duration=30,
                      tracks={"participants": "piste_participants.m4a", "moi": "piste_moi.m4a"})
    session.slides = [SlideRecord(1, 0.0, "d1.png"), SlideRecord(2, 10.0, "d2.png")]
    session.timeline = [(0.0, 1), (10.0, 2)]
    srt = transcribe_session(session, whisper=FakeWhisper(), use_llm=False)

    assert srt.name == "transcription.srt"
    assert Session.load(tmp_path).transcript == "transcription.srt"
    parsed = load_transcript(srt)
    assert [s.speaker for s in parsed] == ["Participants", "Moi", "Participants"]
    assert "[00:00:06] Moi : Merci" in (tmp_path / "transcription.txt").read_text(encoding="utf-8")

    # Sans fichier précisé, le compte-rendu prend la transcription du logiciel, sans décalage,
    # même si un décalage Plaud traîne dans les réglages.
    report = make_report(session, offset=300)
    text = "\n".join(p.text for p in docx.Document(str(report)).paragraphs)
    assert text.index("Merci, j'ai une question") < text.index("Diapo 2")
    assert text.index("Diapo 2") < text.index("Voici les résultats")


@pytest.mark.real_decode
def test_load_audio_with_ffmpeg(tmp_path):
    try:
        ffmpeg = find_ffmpeg()
    except FileNotFoundError:
        pytest.skip("ffmpeg absent")
    path = tmp_path / "piste.m4a"
    subprocess.run([ffmpeg, "-loglevel", "error", "-y", "-f", "lavfi", "-i",
                    "sine=frequency=440:duration=2", "-ac", "2", "-ar", "48000", str(path)], check=True)
    audio = tr_mod.load_audio(path)
    assert audio.dtype == np.float32
    assert abs(audio.size / tr_mod.SAMPLE_RATE - 2.0) < 0.1
    assert 0.05 < np.abs(audio).max() <= 1.0  # sinus lavfi : amplitude 1/8


class FakeTranslator:
    """Traduction factice : préfixe chaque phrase."""

    def translate_sentences(self, sentences):
        return [f"FR<{x}>" for x in sentences]

    def translate(self, texts, progress=None):
        from capture_reunion.summarize import split_sentences

        out = [" ".join(self.translate_sentences(split_sentences(t))) for t in texts]
        if progress:
            progress(1.0)
        return out


ENGLISH = {
    "piste_participants.m4a": [
        (0.0, 6.0, " of the disease. It will be a disease that is less swift in proliferation, but it does not mean the"),
        (6.0, 13.0, " extent. This is really where we look at the stage. So although being sometimes quite impressive when"),
        (13.0, 19.0, " you look at the scans, these low-grade neuroendocrine tumours that are well differentiated often have a"),
        (19.0, 25.0, " multi-year metastatic course, meaning that we are really looking at how to extend the lifespan of"),
        (25.0, 31.0, " our patients, sort of thinking about sequencing different treatments and trying to also look"),
        (31.0, 33.0, " at quality of life while treating our patients."),
        (33.0, 38.0, " And I think there, there is also sometimes a role for liver-directed therapy in combination"),
        (38.0, 45.0, " with other therapies or as a sort of in-between between different other treatment modalities."),
        (45.0, 51.0, " So what we also know is that the vast majority of neuroendocrine tumors have the bulk of"),
        (51.0, 53.0, " their tumor inside the liver."),
    ],
}


class EnglishWhisper(FakeWhisper):
    def transcribe(self, audio, **kwargs):
        self.calls.append((audio.name, kwargs))
        segs = [SimpleNamespace(start=a, end=b, text=t) for a, b, t in ENGLISH[audio.name]]
        return iter(segs), SimpleNamespace(duration=60.0, language="en")


def test_english_talk_paragraphs_summary_and_french_version(tmp_path):
    (tmp_path / "piste_participants.m4a").write_bytes(b"")
    session = Session(folder=tmp_path, started_at="2026-10-03T21:33:00", duration=62,
                      tracks={"participants": "piste_participants.m4a"},
                      window_title="mNETs Webinar")
    session.slides = [SlideRecord(1, 1.0, "d1.png"), SlideRecord(2, 45.0, "d2.png")]
    session.timeline = [(1.0, 1), (45.0, 2)]
    transcribe_session(session, whisper=EnglishWhisper(), translate=True, translator=FakeTranslator(),
                       use_llm=False)

    loaded = Session.load(tmp_path)
    assert loaded.language == "en"
    assert loaded.translations["fr"]["transcript"] == "transcription_fr.srt"
    assert loaded.summary["sentences"]

    txt = (tmp_path / "transcription.txt").read_text(encoding="utf-8")
    blocks = [b for b in txt.split("\n\n") if b.startswith("[")]
    # 10 morceaux de ~6 s → 2 paragraphes (coupure au changement de diapo à 45 s).
    assert len(blocks) == 2
    assert blocks[1].startswith("[00:00:45] Participants : So what we also know")
    assert "SUMMARY" in txt and "Keywords" in txt
    fr_txt = (tmp_path / "transcription_fr.txt").read_text(encoding="utf-8")
    assert "FR<" in fr_txt and "RÉSUMÉ" in fr_txt

    reports = make_reports(loaded)
    assert [p.name for p in reports] == ["compte-rendu.docx", "compte-rendu_fr.docx"]
    en = "\n".join(p.text for p in docx.Document(str(reports[0])).paragraphs)
    fr = "\n".join(p.text for p in docx.Document(str(reports[1])).paragraphs)
    assert "Meeting report" in en and "Slide 1" in en and "Summary" in en
    assert "Before the first slide" not in en  # la 1ʳᵉ diapo apparaît à 1 s
    assert "Compte-rendu de réunion" in fr and "Diapo 2" in fr and "Résumé" in fr and "FR<" in fr
    assert en.count("Participants :") == 2
