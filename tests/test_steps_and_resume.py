"""Diapos dévoilées par étapes, lecture par paquets avec reprise, annulation."""

import json

import cv2
import pytest
from pptx import Presentation

from capture_reunion import rebuild
from capture_reunion.session import Session, SlideRecord
from capture_reunion.transcribe import Cancelled
from capture_reunion.workflow import rebuild_presentation
from slide_fixtures import BULLETS, chart_slide


def _session(tmp_path, images):
    (tmp_path / "diapos").mkdir()
    records = []
    for i, image in enumerate(images, start=1):
        ok, buf = cv2.imencode(".png", image)
        buf.tofile(str(tmp_path / "diapos" / f"diapo_{i:03d}.png"))
        records.append(SlideRecord(i, 10.0 * i, f"diapos/diapo_{i:03d}.png"))
    s = Session(folder=tmp_path, started_at="2026-10-10T10:00:00", duration=10.0 * len(images) + 10)
    s.slides = records
    s.timeline = [(r.timestamp, r.index) for r in records]
    s.save()
    return s


def _steps():
    """Une diapo dévoilée en 3 étapes, puis une autre diapo."""
    return [chart_slide(bullets=BULLETS[:2], chart=False), chart_slide(bullets=BULLETS[:3], chart=False),
            chart_slide(bullets=BULLETS, chart=False),
            chart_slide(title="Liver metastases: options", chart=False)]


def test_progressive_slides_are_grouped(tmp_path):
    s = _session(tmp_path, _steps())
    contents, _ = rebuild.read_slides(s.slide_paths(), s.folder)
    assert rebuild.group_progressive(contents, s.slide_paths()) == [[0, 1, 2], [3]]


def test_same_title_with_new_bullets_is_not_merged(tmp_path):
    other = (("Everolimus in progressive disease", 0), ("Sunitinib for pancreatic NET", 0))
    s = _session(tmp_path, [chart_slide(chart=False), chart_slide(bullets=other, chart=False)])
    contents, _ = rebuild.read_slides(s.slide_paths(), s.folder)
    assert rebuild.group_progressive(contents, s.slide_paths()) == [[0], [1]]


def test_rebuild_keeps_most_complete_step(tmp_path):
    s = _session(tmp_path, _steps())
    out, warnings = rebuild_presentation(s, hybrid_ocr=False)
    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    assert any("2 diapo(s) dévoilée(s) par étapes" in w for w in warnings)
    body = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "Liver-directed therapy" in body  # la dernière étape, complète
    assert "dévoilée en 3 étapes (captures 1, 2, 3)" in prs.slides[0].notes_slide.notes_text_frame.text
    # Avec l'option désactivée, chaque capture reste une diapo.
    out, _ = rebuild_presentation(s, merge_steps=False)
    assert len(Presentation(str(out)).slides) == 4


class CountingEngine:
    """Faux moteur de vision : renvoie un titre par diapo, et compte les diapos lues."""

    label = "LM Studio"
    display_name = "faux modèle"

    def __init__(self, cancel_after=None, skip=()):
        self.read = []
        self.cancel_after = cancel_after
        self.skip = set(skip)

    def connect(self):
        return self.display_name

    def read_slides(self, images):
        out = []
        for k, path in enumerate(images, start=1):
            self.read.append(path.name)
            if path.name not in self.skip:
                out.append({"index": k, "title": f"Titre {path.stem}", "bullets": ["Point"]})
        return out

    def cancelled(self):
        return self.cancel_after is not None and len(self.read) >= self.cancel_after


def _distinct(n):
    return [chart_slide(title=f"Topic number {i}", bullets=((f"Point {i}", 0),), chart=False)
            for i in range(n)]


def test_reading_in_packets_resumes_after_cancel(tmp_path):
    s = _session(tmp_path, _distinct(7))
    engine = CountingEngine(cancel_after=5)
    with pytest.raises(Cancelled):
        rebuild.read_slides(s.slide_paths(), s.folder, engine, cancelled=engine.cancelled)
    assert len(engine.read) == rebuild.PACKET == 5  # un paquet lu, puis arrêt
    cache = json.loads((s.folder / rebuild.CACHE_FILE).read_text(encoding="utf-8"))
    assert sum(v["engine"] == "faux modèle" for v in cache.values()) == 5

    again = CountingEngine()
    contents, warnings = rebuild.read_slides(s.slide_paths(), s.folder, again)
    assert again.read == ["diapo_006.png", "diapo_007.png"]  # reprise là où on s'était arrêté
    assert [c.title for c in contents][-1] == "Titre diapo_007" and warnings == []


def test_unreadable_slide_alone_falls_back_to_ocr(tmp_path):
    s = _session(tmp_path, _distinct(3))
    engine = CountingEngine(skip={"diapo_002.png"})
    contents, warnings = rebuild.read_slides(s.slide_paths(), s.folder, engine)
    assert [c.engine for c in contents] == ["faux modèle", rebuild.OCR_ENGINE, "faux modèle"]
    assert contents[1].title == "Topic number 1"
    assert "1 diapo(s) non lue(s)" in warnings[0]


def test_summaries_resume_after_interruption():
    from capture_reunion.llm import Engine, LLMError

    class Local(Engine):
        label = "LM Studio"

        def __init__(self, fail_on=None):
            super().__init__()
            self.model, self.asked, self.fail_on = "gemma", [], fail_on

        def connect(self):
            return self.model

        def chat(self, system, user, max_tokens=4096, temperature=0.2):
            if "Passage (" in user:
                label = user.split("Passage (")[1].split(")")[0]
                if label == self.fail_on:
                    raise LLMError("LM Studio s'est arrêté")
                self.asked.append(label)
                return f"Résumé {label}"
            return "## Sujet\n- global"

    parts = [(f"Diapo {i}", f"texte {i}") for i in range(1, 8)]
    cache, saves = {}, []
    with pytest.raises(LLMError):
        Local(fail_on="Diapo 6").summarize_all(parts, "fr", cache=cache, on_part=lambda: saves.append(1))
    assert len(cache) == len(saves) == 5
    again = Local()
    summarized, overall = again.summarize_all(parts, "fr", cache=cache)
    assert again.asked == ["Diapo 6", "Diapo 7"] and summarized[0] == ("Diapo 1", "Résumé Diapo 1")
    assert overall.startswith("## Sujet")
