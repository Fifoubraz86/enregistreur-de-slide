"""Client LM Studio testé contre un faux serveur HTTP local qui imite son API."""

import json
import re
import threading
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from types import SimpleNamespace

import docx
import numpy as np
import pytest
from pptx import Presentation

from capture_reunion import transcribe as tr_mod
from capture_reunion import translate as translate_mod
from capture_reunion.llm import LLMError, LMStudio, _parse_markers, load_glossary
from capture_reunion.session import Session, SlideRecord
from capture_reunion.transcribe import transcribe_session
from capture_reunion.workflow import make_pptx, make_reports


def fr(text: str) -> str:
    """Fausse traduction, mais bien en français : reprend les 3 premiers mots pour les tests."""
    return f"Traduction française de « {' '.join(text.split()[:3])} » pour les médecins."


class FakeLMStudio(BaseHTTPRequestHandler):
    calls: list = []
    drop_marker = False  # simule un modèle qui « oublie » un paragraphe
    style = "brackets"  # brackets | bold | none | english
    jit = True  # chargement « Just-In-Time » activé dans LM Studio

    def log_message(self, *args):
        pass

    def _send(self, payload):
        body = json.dumps(payload).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        if self.path == "/api/v0/models":
            self._send({"data": [
                {"id": "text-embedding-nomic", "type": "embeddings", "state": "loaded"},
                {"id": "qwen-autre", "type": "llm", "state": "not-loaded"},
                {"id": "google/gemma-4-26b-a4b-qat", "type": "llm", "state": "loaded"},
            ]})
        else:
            self.send_error(404)

    def do_POST(self):
        req = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        if req["model"] == "qwen-autre" and not FakeLMStudio.jit:
            body = b'{"error": "Model is not loaded"}'
            self.send_response(400)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        system, user = req["messages"][0]["content"], req["messages"][1]["content"]
        FakeLMStudio.calls.append((req["model"], system, user))
        if "marqueur [n]" in system:  # traduction par blocs
            items = re.findall(r"^\[(\d+)\] (.*)$", user, re.M)
            if FakeLMStudio.drop_marker and len(items) > 1:
                items = items[:-1]
            style = FakeLMStudio.style
            if style == "bold":
                lines = [f"**[{n}]** {fr(t)}" for n, t in items]
            elif style == "none":
                lines = [fr(t) for _, t in items]
            elif style == "english":
                lines = [f"[{n}] {t}" for n, t in items]
            else:
                lines = [f"[{n}] {fr(t)}" for n, t in items]
            content = "<think>je réfléchis</think>\n```\n" + "\n\n".join(lines) + "\n```"
        elif "mise en forme" in system:
            if FakeLMStudio.style == "english":
                content = user
            elif user.lstrip().startswith("##"):
                content = "\n".join(line if line.startswith("##") else "- " + fr(line)
                                     for line in user.splitlines() if line.strip())
            else:
                content = fr(user)
        elif "Rédige un résumé structuré" in user:
            content = "## Sujet\n- TNE hépatiques\n## Messages clés\n- **Survie** prolongée"
        else:
            label = re.search(r"Passage \((.+?)\)", user).group(1)
            content = f"Résumé de {label}."
        self._send({"choices": [{"message": {"content": content}}]})


@pytest.fixture
def server():
    FakeLMStudio.calls = []
    FakeLMStudio.drop_marker = False
    FakeLMStudio.style = "brackets"
    FakeLMStudio.jit = True
    httpd = HTTPServer(("127.0.0.1", 0), FakeLMStudio)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}"
    httpd.shutdown()


def test_connect_prefers_loaded_llm(server):
    client = LMStudio(server + "/v1")
    assert client.connect() == "google/gemma-4-26b-a4b-qat"


def test_model_infos_mark_loaded(server):
    infos = LMStudio(server).model_infos()
    assert infos == [{"id": "google/gemma-4-26b-a4b-qat", "loaded": True},
                     {"id": "qwen-autre", "loaded": False}]


def test_chosen_model_is_used_even_if_another_is_loaded(server):
    """Régression : le modèle choisi est utilisé (chargé à la demande), pas celui déjà chargé."""
    client = LMStudio(server, model="qwen-autre")
    assert client.connect() == "qwen-autre" and client.needs_loading
    client.chat("s", "Passage (x) : bonjour")
    assert FakeLMStudio.calls[-1][0] == "qwen-autre"
    assert not client.needs_loading


def test_chosen_model_load_failure_is_explained(server):
    FakeLMStudio.jit = False
    client = LMStudio(server, model="qwen-autre")
    client.connect()
    with pytest.raises(LLMError, match="Just-In-Time"):
        client.chat("s", "Passage (x) : bonjour")


def test_unknown_model_is_an_error(server):
    with pytest.raises(LLMError, match="n'existe pas"):
        LMStudio(server, model="gemma-inexistant").connect()


def test_pipeline_uses_chosen_model(server, session):
    transcribe_session(session, whisper=Whisper(), llm_url=server, llm_model="qwen-autre")
    assert Session.load(session.folder).ai["model"] == "qwen-autre"
    assert {c[0] for c in FakeLMStudio.calls} == {"qwen-autre"}


def test_unreachable_server_message():
    with pytest.raises(LLMError, match="LM Studio injoignable"):
        LMStudio("http://127.0.0.1:9").connect()


TEXTS = [f"Paragraph {i} is about the liver and the tumours of our patients. " + "x" * 600
         for i in range(5)]


def test_translate_batches_keep_order_and_recover_lost_paragraphs(server):
    FakeLMStudio.drop_marker = True
    client = LMStudio(server)
    out = client.translate(TEXTS, "en", "fr", glossary="NET = TNE", batch_chars=1500)
    assert out == [fr(t) for t in TEXTS]
    assert client.last_failures == 0
    systems = [c[1] for c in FakeLMStudio.calls]
    assert all("NET = TNE" in s for s in systems)
    assert len(FakeLMStudio.calls) > 3  # blocs + reprises unitaires


@pytest.mark.parametrize("style", ["bold", "none"])
def test_translate_tolerates_marker_styles(server, style):
    """Régression : numéros en gras ou absents ne doivent pas laisser le texte en anglais."""
    FakeLMStudio.style = style
    client = LMStudio(server)
    assert client.translate(TEXTS[:3], "en", "fr") == [fr(t) for t in TEXTS[:3]]
    assert client.last_failures == 0


def test_untranslated_output_is_reported(server):
    FakeLMStudio.style = "english"
    client = LMStudio(server)
    out = client.translate(TEXTS[:2], "en", "fr")
    assert out == TEXTS[:2] and client.last_failures == 2


def test_looks_untranslated():
    from capture_reunion.llm import looks_untranslated

    en = "So what we also know is that the vast majority of tumours have the bulk in the liver."
    assert looks_untranslated(en, en, "en", "fr")
    assert looks_untranslated(en, en.upper().lower(), "en", "fr")
    assert not looks_untranslated(en, "Nous savons aussi que la grande majorité des tumeurs sont dans le foie.",
                                  "en", "fr")


def test_parse_markers_multiline():
    assert _parse_markers("[1] a\nsuite\n[2] b") == {1: "a suite", 2: "b"}


def test_glossary(tmp_path):
    g = tmp_path / "glossaire.txt"
    g.write_text("# commentaire\nNET = TNE\n\nPRRT = RIV\n", encoding="utf-8")
    assert load_glossary(g) == "NET = TNE\nPRRT = RIV"
    assert load_glossary(None) == ""


SCRIPT = [(0.0, 6.0, " Neuroendocrine tumours have a long course."),
          (6.0, 12.0, " We extend the lifespan of our patients."),
          (45.0, 51.0, " The bulk of the tumour is inside the liver.")]


class Whisper:
    def transcribe(self, audio, **kwargs):
        return iter([SimpleNamespace(start=a, end=b, text=t) for a, b, t in SCRIPT]), \
            SimpleNamespace(language="en")


@pytest.fixture
def session(tmp_path, monkeypatch):
    monkeypatch.setattr(tr_mod, "load_audio", lambda p: np.zeros(60 * 16000, np.float32))
    (tmp_path / "piste_participants.m4a").write_bytes(b"")
    s = Session(folder=tmp_path, started_at="2026-10-04T10:00:00", duration=60,
                tracks={"participants": "piste_participants.m4a"}, window_title="mNETs Webinar")
    import cv2
    for i in (1, 2):
        cv2.imwrite(str(tmp_path / f"d{i}.png"), np.full((90, 160, 3), 200 * (i - 1), np.uint8))
    s.slides = [SlideRecord(1, 0.0, "d1.png"), SlideRecord(2, 40.0, "d2.png")]
    s.timeline = [(0.0, 1), (40.0, 2)]
    return s


def test_full_pipeline_with_lmstudio(server, session):
    transcribe_session(session, whisper=Whisper(), translate=True, llm_url=server, glossary="NET = TNE")
    loaded = Session.load(session.folder)
    assert loaded.ai["model"] == "google/gemma-4-26b-a4b-qat"
    assert loaded.ai["slide_summaries"]["en"] == {"1": "Résumé de Slide 1.", "2": "Résumé de Slide 2."}
    assert loaded.ai["summary"]["en"].startswith("## Sujet")
    assert loaded.ai["summary"]["fr"].startswith("## Sujet")
    assert loaded.translations["fr"]["engine"] == "IA locale (google/gemma-4-26b-a4b-qat)"
    assert loaded.ai["warnings"] == []

    fr_txt = (session.folder / "transcription_fr.txt").read_text(encoding="utf-8")
    assert "Traduction française de « Neuroendocrine tumours have »" in fr_txt
    assert "We extend" not in fr_txt  # le transcrit lui-même est traduit, pas seulement le résumé
    assert "IA locale google/gemma-4-26b-a4b-qat" in fr_txt and "PHRASES CLÉS" in fr_txt
    assert "<think>" not in fr_txt

    en_doc, fr_doc = make_reports(loaded)
    en = "\n".join(p.text for p in docx.Document(str(en_doc)).paragraphs)
    fr = "\n".join(p.text for p in docx.Document(str(fr_doc)).paragraphs)
    assert "In brief: Résumé de Slide 1." in en
    assert "Messages clés" in en and "Survie prolongée" in en  # rendu des « ## » et du gras
    assert "Key sentences actually spoken" in en
    assert "En bref : Traduction française de « Résumé de Slide »" in fr
    assert "Phrases clés réellement prononcées" in fr
    assert "The bulk of the tumour" not in fr

    pptx = make_pptx(loaded)
    notes = Presentation(str(pptx)).slides[0].notes_slide.notes_text_frame.text
    assert notes.startswith("In brief: Résumé de Slide 1.")


def test_without_lmstudio_falls_back_to_argos(session):
    class Argos:
        def translate(self, texts, progress=None):
            return [f"ARGOS({t})" for t in texts]

        def translate_sentences(self, sentences):
            return [f"ARGOS({x})" for x in sentences]

    transcribe_session(session, whisper=Whisper(), translate=True, translator=Argos(),
                       llm_url="http://127.0.0.1:9")
    loaded = Session.load(session.folder)
    assert loaded.translations["fr"]["engine"] == "Argos"
    assert "ARGOS(" in (session.folder / "transcription_fr.txt").read_text(encoding="utf-8")
    assert "IA locale non utilisée" in loaded.ai["warnings"][0]
    assert "summary" not in loaded.ai


def test_translation_failure_keeps_transcript(session):
    class Broken:
        def translate(self, *a, **k):
            raise RuntimeError("HTTP Error 403: Forbidden")

    transcribe_session(session, whisper=Whisper(), translate=True, translator=Broken(), use_llm=False)
    loaded = Session.load(session.folder)
    assert (session.folder / "transcription.txt").exists() and loaded.translations == {}
    assert "Traduction impossible : HTTP Error 403" in loaded.ai["warnings"][0]


def test_argos_download_sends_browser_user_agent(tmp_path, monkeypatch):
    seen = []

    def fake_urlopen(request, timeout=0):
        seen.append(request.get_header("User-agent"))
        raise OSError("HTTP Error 403: Forbidden")

    monkeypatch.setattr(translate_mod.urllib.request, "urlopen", fake_urlopen)
    with pytest.raises(RuntimeError) as err:
        translate_mod.ensure_model("en", "fr", root=tmp_path)
    assert seen and seen[0].startswith("Mozilla/5.0")
    message = str(err.value)
    assert "translate-en_fr-1_9.argosmodel" in message and str(tmp_path) in message


def test_argos_manually_dropped_model_is_used(tmp_path):
    root = tmp_path / "traduction-en_fr"
    root.mkdir()
    with zipfile.ZipFile(root / "translate-en_fr-1_9.argosmodel", "w") as z:
        z.writestr("translate-en_fr-1_9/model/model.bin", b"x")
        z.writestr("translate-en_fr-1_9/sentencepiece.model", b"x")
    found = translate_mod.ensure_model("en", "fr", root=tmp_path)
    assert found.name == "translate-en_fr-1_9"
    assert not list(root.glob("*.argosmodel"))
