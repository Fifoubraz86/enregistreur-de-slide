import copy
import zipfile

import cv2
import docx
from pptx import Presentation
from pptx.util import Inches

from capture_reunion.export import build_pptx, build_report
from capture_reunion.transcript import assign_segments, parse_transcript
from conftest import make_slide


def _images(tmp_path, n=3):
    paths = []
    for i in range(1, n + 1):
        p = tmp_path / f"diapo_{i}.png"
        cv2.imwrite(str(p), make_slide(i))
        paths.append(p)
    return paths


def _make_template(tmp_path):
    """Modèle .potx avec un bandeau sur le masque (comme un logo CHU)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    example = prs.slides.add_slide(prs.slide_layouts[0])  # diapo existante à retirer
    banner = example.shapes.add_textbox(0, 0, prs.slide_width, Inches(0.5))
    banner.text_frame.text = "BANDEAU MAISON"
    prs.slide_master.shapes._spTree.append(copy.deepcopy(banner._element))
    pptx_path = tmp_path / "modele.pptx"
    prs.save(str(pptx_path))
    potx_path = tmp_path / "modele.potx"
    with zipfile.ZipFile(pptx_path) as src, zipfile.ZipFile(potx_path, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentation.main+xml", b"template.main+xml")
            dst.writestr(item, data)
    return potx_path


def test_pptx_without_template(tmp_path):
    out = build_pptx(_images(tmp_path), tmp_path / "out.pptx", notes={2: "note deux"})
    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    assert abs(prs.slide_height / prs.slide_width - 360 / 640) < 0.01
    assert prs.slides[1].notes_slide.notes_text_frame.text == "note deux"
    pics = [s for s in prs.slides[0].shapes if s.shape_type == 13]
    assert len(pics) == 1


def test_pptx_with_potx_template(tmp_path):
    out = build_pptx(_images(tmp_path), tmp_path / "out.pptx", template=_make_template(tmp_path))
    prs = Presentation(str(out))
    assert len(prs.slides) == 3  # la diapo d'exemple du modèle a été retirée
    texts = [sh.text_frame.text for sh in prs.slide_master.shapes if sh.has_text_frame]
    assert "BANDEAU MAISON" in texts
    for slide in prs.slides:
        # Seul le numéro de diapo reste (s'il existe dans la disposition).
        assert all("SLIDE_NUMBER" in str(ph.placeholder_format.type) for ph in slide.placeholders)
        pic = slide.shapes[0]
        assert pic.left >= 0 and pic.left + pic.width <= prs.slide_width


def test_report(tmp_path):
    images = _images(tmp_path, 2)
    segs = parse_transcript("[00:00:02] A: un\n[00:00:12] B: deux\n")
    sections = assign_segments([(0.0, 1), (10.0, 2)], segs, duration=20)
    out = build_report(sections, {1: images[0], 2: images[1]}, tmp_path / "cr.docx")
    d = docx.Document(str(out))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Diapo 1" in text and "Diapo 2" in text and "B : deux" in text
    assert len(d.inline_shapes) == 2
