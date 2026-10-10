"""Reconstruction des diapos : à partir des captures, une vraie présentation PowerPoint
avec du texte modifiable (titre, puces) et les figures découpées, dans un masque au
choix (celui de l'utilisateur, ou le masque par défaut).

Lecture des diapos :
1. par une IA qui voit les images (Claude via Claude Code, ou un modèle de vision dans
   LM Studio) : meilleure structure (titre, niveaux de puces, figures) ;
2. sinon par OCR local (RapidOCR, hors ligne) : texte fiable, structure déduite de la
   mise en page.

Les lectures sont gardées dans ``diapos_lues.json`` : refaire le PowerPoint avec un
autre masque ne relit pas les diapos (et ne reconsomme pas de quota).
"""

from __future__ import annotations

import copy
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Callable, Optional

import cv2
import numpy as np

CACHE_FILE = "diapos_lues.json"
FIGURES_DIR = "diapos_figures"
_BULLET = re.compile(r"^\s*(?:[•●○◦▪▫■□‣∙·\-–—*>]|\d{1,2}[.)])\s*")

Box = tuple[float, float, float, float]  # x, y, largeur, hauteur en fractions


@dataclass
class SlideContent:
    index: int
    title: str = ""
    bullets: list[tuple[str, int]] = field(default_factory=list)
    figures: list[Box] = field(default_factory=list)
    source: str = ""
    """Image d'origine (chemin relatif au dossier de session)."""
    engine: str = ""

    @classmethod
    def from_dict(cls, data: dict) -> "SlideContent":
        bullets = []
        for b in data.get("bullets") or []:
            if isinstance(b, dict):
                text, level = str(b.get("text", "")).strip(), int(b.get("level") or 0)
            else:
                text, level = str(b).strip(), 0
            if text:
                bullets.append((_BULLET.sub("", text).strip() or text, max(0, min(2, level))))
        figures = []
        for f in data.get("figures") or []:
            box = f.get("box") if isinstance(f, dict) else f
            if isinstance(box, (list, tuple)) and len(box) == 4:
                figures.append(_clamp_box(tuple(float(v) for v in box)))
        return cls(int(data.get("index") or 0), str(data.get("title") or "").strip(), bullets,
                   [f for f in figures if f[2] > 0.02 and f[3] > 0.02],
                   str(data.get("source") or ""), str(data.get("engine") or ""))


def _clamp_box(box) -> Box:
    x, y, w, h = box
    x, y = max(0.0, min(1.0, x)), max(0.0, min(1.0, y))
    return x, y, max(0.0, min(1.0 - x, w)), max(0.0, min(1.0 - y, h))


def read_image(path: Path) -> np.ndarray:
    data = np.fromfile(str(path), dtype=np.uint8)
    image = cv2.imdecode(data, cv2.IMREAD_COLOR)
    if image is None:
        raise ValueError(f"Image illisible : {path}")
    return image


# ---------------------------------------------------------------- OCR local
_ocr_engine = None


def _ocr(image: np.ndarray) -> list[tuple[np.ndarray, str, float]]:
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError as exc:
            raise RuntimeError("Module d'OCR absent : pip install rapidocr_onnxruntime") from exc
        _ocr_engine = RapidOCR()
    result, _ = _ocr_engine(image)
    return [(np.array(box, dtype=np.float32), str(text), float(score)) for box, text, score in (result or [])]


@dataclass
class _Line:
    x0: float
    y0: float
    x1: float
    y1: float
    text: str

    @property
    def h(self) -> float:
        return self.y1 - self.y0

    @property
    def cy(self) -> float:
        return (self.y0 + self.y1) / 2


def detect_figures(image: np.ndarray, text_boxes: list[tuple[float, float, float, float]]) -> list[Box]:
    """Zones de graphiques, photos, schémas : ce qui n'est ni du texte ni le fond."""
    h, w = image.shape[:2]
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    # Fond = couleur la plus fréquente ; contenu = ce qui s'en écarte, ou les contours.
    small = cv2.resize(image, (64, 36)).reshape(-1, 3)
    values, counts = np.unique((small // 16), axis=0, return_counts=True)
    background = values[counts.argmax()] * 16 + 8
    diff = np.abs(image.astype(np.int16) - background.astype(np.int16)).max(axis=2) > 40
    edges = cv2.Canny(gray, 60, 160) > 0
    mask = (diff | edges).astype(np.uint8)
    pad = max(2, h // 150)
    for x0, y0, x1, y1 in text_boxes:
        cv2.rectangle(mask, (int(x0) - pad, int(y0) - pad), (int(x1) + pad, int(y1) + pad), 0, -1)
    k = max(5, w // 90)
    mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, np.ones((k, k), np.uint8))
    mask = cv2.dilate(mask, np.ones((k // 2 + 1, k // 2 + 1), np.uint8))
    count, _, stats, _ = cv2.connectedComponentsWithStats(mask)
    boxes = []
    for i in range(1, count):
        x, y, bw, bh, area = stats[i]
        frac = bw * bh / (w * h)
        if frac < 0.03 or frac > 0.92:
            continue
        if bw / max(1, bh) > 8 or bh < h * 0.08:  # bandeaux et filets décoratifs
            continue
        if area / (bw * bh) < 0.15:  # cadre vide
            continue
        boxes.append((x / w, y / h, bw / w, bh / h))
    return _merge_boxes(boxes)


def _iou(a: Box, b: Box) -> float:
    ax1, ay1, bx1, by1 = a[0] + a[2], a[1] + a[3], b[0] + b[2], b[1] + b[3]
    iw = max(0.0, min(ax1, bx1) - max(a[0], b[0]))
    ih = max(0.0, min(ay1, by1) - max(a[1], b[1]))
    inter = iw * ih
    union = a[2] * a[3] + b[2] * b[3] - inter
    return inter / union if union else 0.0


def _merge_boxes(boxes: list[Box]) -> list[Box]:
    boxes = list(boxes)
    merged = True
    while merged:
        merged = False
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                if _iou(a, b) > 0 or _contains(a, b) or _contains(b, a):
                    x0, y0 = min(a[0], b[0]), min(a[1], b[1])
                    x1, y1 = max(a[0] + a[2], b[0] + b[2]), max(a[1] + a[3], b[1] + b[3])
                    boxes[i] = (x0, y0, x1 - x0, y1 - y0)
                    del boxes[j]
                    merged = True
                    break
            if merged:
                break
    return sorted(boxes, key=lambda b: (b[1], b[0]))


def _contains(a: Box, b: Box) -> bool:
    return a[0] <= b[0] and a[1] <= b[1] and a[0] + a[2] >= b[0] + b[2] and a[1] + a[3] >= b[1] + b[3]


def ocr_slide(image: np.ndarray, index: int = 0) -> SlideContent:
    """Lecture hors ligne : titre, puces (avec niveaux) et figures, déduits de la mise en page."""
    h, w = image.shape[:2]
    lines = []
    for box, text, score in _ocr(image):
        if score < 0.5 or not text.strip():
            continue
        lines.append(_Line(float(box[:, 0].min()), float(box[:, 1].min()),
                           float(box[:, 0].max()), float(box[:, 1].max()), text.strip()))
    figures = detect_figures(image, [(l.x0, l.y0, l.x1, l.y1) for l in lines])
    # Les petits textes collés à une figure (légendes d'axes, sources) en font partie.
    near = 0.03
    for k, f in enumerate(figures):
        x0, y0, x1, y1 = f[0] - near, f[1] - near, f[0] + f[2] + near, f[1] + f[3] + near
        for line in lines:
            cx, cy = (line.x0 + line.x1) / 2 / w, line.cy / h
            if x0 <= cx <= x1 and y0 <= cy <= y1:
                fx0, fy0 = min(f[0], line.x0 / w), min(f[1], line.y0 / h)
                fx1, fy1 = max(f[0] + f[2], line.x1 / w), max(f[1] + f[3], line.y1 / h)
                f = (fx0, fy0, fx1 - fx0, fy1 - fy0)
        figures[k] = _clamp_box(f)

    def in_figure(line: _Line) -> bool:
        cx, cy = (line.x0 + line.x1) / 2 / w, line.cy / h
        return any(f[0] <= cx <= f[0] + f[2] and f[1] <= cy <= f[1] + f[3] for f in figures)

    text_lines = [l for l in lines if not in_figure(l)]  # légendes d'axes, etc. : restent dans l'image
    if not text_lines:
        return SlideContent(index, figures=figures, engine="OCR local")
    text_lines.sort(key=lambda l: (l.cy, l.x0))

    # Titre : la plus grande écriture dans le haut de la diapo.
    top = [l for l in text_lines if l.y0 < h * 0.3]
    median_h = float(np.median([l.h for l in text_lines]))
    title_lines: list[_Line] = []
    if top:
        head = max(top, key=lambda l: l.h)
        if head.h >= median_h * 1.15 or head.y0 < h * 0.18:
            # Le titre : la plus grande ligne, plus celles de même taille qui la suivent
            # immédiatement (titre sur deux lignes).
            title_lines = [head]
            bottom = head.y1
            for line in sorted(top, key=lambda l: l.cy):
                if line is head or line.h < head.h * 0.85:
                    continue
                if abs(line.cy - head.cy) < head.h * 0.5 or 0 <= line.y0 - bottom < head.h * 0.6:
                    title_lines.append(line)
                    bottom = max(bottom, line.y1)
    title = " ".join(l.text for l in sorted(title_lines, key=lambda l: (l.cy, l.x0)))
    body = [l for l in text_lines if l not in title_lines]

    # Lignes d'une même rangée, puis paragraphes (puces) et niveaux d'après le retrait.
    rows: list[list[_Line]] = []
    for line in body:
        if rows and abs(rows[-1][0].cy - line.cy) < line.h * 0.5:
            rows[-1].append(line)
        else:
            rows.append([line])
    bullets: list[tuple[str, int]] = []
    if rows:
        min_x = min(r[0].x0 for r in rows)
        unit = max(median_h * 1.5, w * 0.02)
        prev = None
        for row in rows:
            row.sort(key=lambda l: l.x0)
            text = " ".join(l.text for l in row)
            x0, y0, rh = row[0].x0, min(l.y0 for l in row), max(l.h for l in row)
            level = max(0, min(2, int(round((x0 - min_x) / unit))))
            starts_bullet = bool(_BULLET.match(text))
            continues = (prev is not None and not starts_bullet and y0 - prev[1] < rh * 0.9
                         and abs(x0 - prev[0]) < unit * 0.6 and not bullets[-1][0].endswith((".", ":")))
            clean = _BULLET.sub("", text).strip() or text
            if continues:
                bullets[-1] = (f"{bullets[-1][0]} {clean}", bullets[-1][1])
            else:
                bullets.append((clean, level))
            prev = (x0, max(l.y1 for l in row))
    return SlideContent(index, title, bullets, figures, engine="OCR local")


def refine_figures(ai_boxes: list[Box], detected: list[Box]) -> list[Box]:
    """Les rectangles donnés par l'IA sont approximatifs : on les cale sur ceux détectés."""
    out = []
    for box in ai_boxes:
        best = max(detected, key=lambda d: _iou(box, d), default=None)
        out.append(best if best is not None and _iou(box, best) > 0.3 else box)
    return out


# ---------------------------------------------------------------- lecture
def read_slides(
    images: list[Path],
    folder: Path,
    engine=None,
    progress: Optional[Callable[[float], None]] = None,
    use_cache: bool = True,
) -> tuple[list[SlideContent], list[str]]:
    """Lit toutes les diapos (IA si possible, sinon OCR). Renvoie (contenus, avertissements)."""
    from .llm import LLMError

    warnings: list[str] = []
    cache_path = Path(folder) / CACHE_FILE
    cache = {}
    if use_cache and cache_path.exists():
        try:
            cache = json.loads(cache_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            cache = {}

    def key(path: Path) -> str:
        stat = Path(path).stat()
        return f"{Path(path).name}:{stat.st_size}:{int(stat.st_mtime)}"

    contents: dict[int, SlideContent] = {}
    todo = []
    for n, path in enumerate(images, start=1):
        cached = cache.get(key(path))
        if cached and (engine is None or cached.get("engine") != "OCR local"):
            contents[n] = SlideContent.from_dict({**cached, "index": n})
        else:
            todo.append((n, path))

    if todo and engine is not None:
        try:
            engine.connect()
            read = engine.read_slides([p for _, p in todo])
            by_index = {int(d.get("index") or 0): d for d in read}
            for k, (n, path) in enumerate(todo, start=1):
                data = by_index.get(k)
                if data:
                    content = SlideContent.from_dict({**data, "index": n})
                    if content.figures:
                        image = read_image(path)
                        try:
                            text_boxes = [(float(b[:, 0].min()), float(b[:, 1].min()),
                                           float(b[:, 0].max()), float(b[:, 1].max()))
                                          for b, _, _ in _ocr(image)]
                        except RuntimeError:
                            text_boxes = []
                        content.figures = refine_figures(content.figures, detect_figures(image, text_boxes))
                    content.engine = engine.display_name
                    contents[n] = content
        except LLMError as exc:
            warnings.append(f"Lecture par {engine.label} impossible ({exc}) : OCR local utilisé.")
        if progress:
            progress(0.5)

    remaining = [(n, p) for n, p in todo if n not in contents]
    for i, (n, path) in enumerate(remaining, start=1):
        contents[n] = ocr_slide(read_image(path), n)
        if progress:
            progress(0.5 + 0.5 * i / len(remaining))

    for n, path in enumerate(images, start=1):
        c = contents[n]
        c.index = n
        c.source = Path(path).name
        cache[key(path)] = {**asdict(c), "bullets": [{"text": t, "level": lv} for t, lv in c.bullets],
                            "figures": [{"box": list(b)} for b in c.figures]}
    try:
        cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    except OSError:
        pass
    if progress:
        progress(1.0)
    return [contents[n] for n in range(1, len(images) + 1)], warnings


# ---------------------------------------------------------------- PowerPoint
def default_template_path() -> Optional[Path]:
    import sys

    roots = [Path(sys._MEIPASS)] if getattr(sys, "_MEIPASS", None) else []
    roots.append(Path(__file__).resolve().parent.parent)
    for root in roots:
        candidate = root / "modeles" / "masque_par_defaut.pptx"
        if candidate.exists():
            return candidate
    return None


def make_default_template(output: Path) -> Path:
    """Masque sobre 16:9 : bandeau bleu en haut, filet et numéro de page en bas."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    ratio = Inches(13.333) / prs.slide_width
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for container in [prs.slide_master, *prs.slide_layouts]:
        for shape in container.placeholders:
            # Lire les 4 valeurs (éventuellement héritées) avant d'en écrire une : écrire
            # seulement la largeur remettrait la position et la hauteur à zéro.
            if shape._element.spPr.find(qn("a:xfrm")) is None:
                continue  # position héritée du masque, déjà élargie
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            if None in (left, top, width, height):
                continue
            shape.left, shape.top = Emu(int(left * ratio)), Emu(int(top))
            shape.width, shape.height = Emu(int(width * ratio)), Emu(int(height))
    # Les formes du masque se créent sur une diapo temporaire puis sont copiées.
    tmp = prs.slides.add_slide(prs.slide_layouts[6])
    band = tmp.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
    rule = tmp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(7.05),
                                prs.slide_width - Inches(1.0), Emu(12700))
    for shape, color in ((band, RGBColor(0x1F, 0x4E, 0x79)), (rule, RGBColor(0xA6, 0xA6, 0xA6))):
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        prs.slide_master.shapes._spTree.append(copy.deepcopy(shape._element))
    sld_id = prs.slides._sldIdLst[0]
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)
    title_style = prs.slide_master.element.find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}titleStyle")
    if title_style is not None:
        from lxml import etree

        for rpr in title_style.iter(qn("a:defRPr")):
            rpr.set("sz", str(Pt(32).centipoints))
            for old in rpr.findall(qn("a:solidFill")):
                rpr.remove(old)
            fill = etree.SubElement(rpr, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", "1F4E79")
            rpr.insert(0, fill)  # la couleur doit précéder la police dans le schéma
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output


def _find_layout(prs, with_body: bool):
    from pptx.enum.shapes import PP_PLACEHOLDER

    def kinds(layout):
        return {ph.placeholder_format.type for ph in layout.placeholders}

    body_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    candidates = [l for l in prs.slide_layouts if kinds(l) & title_types]
    if with_body:
        named = [l for l in candidates if kinds(l) & body_types
                 and any(k in l.name.lower() for k in ("contenu", "content", "texte", "text"))]
        bodied = [l for l in candidates if kinds(l) & body_types]
        if named or bodied:
            return (named or bodied)[0]
    titled = [l for l in candidates if any(k in l.name.lower() for k in ("titre seul", "title only"))]
    return (titled or candidates or list(prs.slide_layouts))[0]


def _font_size(bullets: list[tuple[str, int]], narrow: bool):
    from pptx.util import Pt

    chars = sum(len(t) for t, _ in bullets) * (1.8 if narrow else 1.0)
    lines = len(bullets) + chars / 90
    return Pt(max(12, min(26, round(30 - 1.6 * lines))))


def build_rebuilt_pptx(
    contents: list[SlideContent],
    folder: Path,
    output: Path,
    template: Optional[Path] = None,
    notes: Optional[dict[int, str]] = None,
) -> Path:
    """Présentation reconstruite : titre et puces en texte modifiable, figures en images."""
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Emu

    from .export import _remove_existing_slides, open_template

    folder = Path(folder)
    template = template or default_template_path()
    if template is None:
        template = make_default_template(folder / FIGURES_DIR / "masque_par_defaut.pptx")
    prs = open_template(template)
    _remove_existing_slides(prs)
    fig_dir = folder / FIGURES_DIR
    fig_dir.mkdir(exist_ok=True)
    margin = int(prs.slide_width * 0.04)

    for c in contents:
        has_body = bool(c.bullets)
        layout = _find_layout(prs, has_body)
        slide = prs.slides.add_slide(layout)
        title_ph = body_ph = None
        for ph in list(slide.placeholders):
            kind = ph.placeholder_format.type
            if kind in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) and title_ph is None:
                title_ph = ph
            elif kind in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT) and body_ph is None and has_body:
                body_ph = ph
            else:
                ph._element.getparent().remove(ph._element)  # dates, pieds de page, cadres vides…

        if title_ph is not None:
            if c.title:
                title_ph.text_frame.text = c.title
            else:
                title_ph._element.getparent().remove(title_ph._element)
                title_ph = None

        # Zone disponible pour le contenu (sous le titre).
        top = (title_ph.top + title_ph.height) if title_ph is not None else margin
        area = (margin, top + margin // 2, prs.slide_width - 2 * margin,
                prs.slide_height - top - int(margin * 1.5))
        if body_ph is not None:
            area = (body_ph.left, body_ph.top, body_ph.width, body_ph.height)

        image = read_image(folder / "diapos" / c.source) if c.figures else None
        if body_ph is not None:
            if c.figures:  # texte à gauche, figures à droite (les 4 valeurs ensemble, cf. masque)
                body_ph.left, body_ph.top = Emu(area[0]), Emu(area[1])
                body_ph.width, body_ph.height = Emu(int(area[2] * 0.52)), Emu(area[3])
            tf = body_ph.text_frame
            size = _font_size(c.bullets, narrow=bool(c.figures))
            for k, (text, level) in enumerate(c.bullets):
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                para.text = text
                para.level = level
                for run in para.runs:
                    run.font.size = size
        if c.figures and image is not None:
            ih, iw = image.shape[:2]
            if body_ph is not None:
                fx = area[0] + int(area[2] * 0.55)
                fbox = (fx, area[1], area[0] + area[2] - fx, area[3])
            else:
                fbox = area
            slot_h = fbox[3] // len(c.figures)
            for k, (x, y, w, h) in enumerate(c.figures):
                crop = image[int(y * ih):int((y + h) * ih), int(x * iw):int((x + w) * iw)]
                if crop.size == 0:
                    continue
                path = fig_dir / f"diapo_{c.index:03d}_figure_{k + 1}.png"
                ok, buf = cv2.imencode(".png", crop)
                if not ok:
                    continue
                buf.tofile(str(path))
                ch, cw = crop.shape[:2]
                scale = min(fbox[2] / cw, (slot_h - margin // 4) / ch)
                pw, phh = int(cw * scale), int(ch * scale)
                left = fbox[0] + (fbox[2] - pw) // 2
                top_k = fbox[1] + k * slot_h + (slot_h - phh) // 2
                slide.shapes.add_picture(str(path), left, top_k, pw, phh)
        note = (notes or {}).get(c.index, "")
        source = f"Capture d'origine : diapos/{c.source} — lecture : {c.engine}"
        slide.notes_slide.notes_text_frame.text = f"{note}\n\n{source}".strip()

    output = Path(output)
    prs.save(str(output))
    return output
