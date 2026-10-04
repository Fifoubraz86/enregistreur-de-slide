"""Export des diapos capturées : PowerPoint (avec votre modèle) et compte-rendu Word."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Optional

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu, Pt

from .slides import format_timestamp
from .transcript import SlideSection

_TEMPLATE_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
_PRESENTATION_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
_BLANK_HINTS = ("vide", "blank", "leer", "vacío")


def open_template(path: Optional[Path]):
    """Ouvre un modèle .pptx ou .potx (python-pptx ne lit pas directement les .potx)."""
    if path is None:
        return Presentation()
    path = Path(path)
    if path.suffix.lower() not in (".potx", ".potm"):
        return Presentation(str(path))
    src = zipfile.ZipFile(path)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(_TEMPLATE_CT.encode(), _PRESENTATION_CT.encode())
            dst.writestr(item, data)
    out.seek(0)
    return Presentation(out)


def _remove_existing_slides(prs) -> None:
    id_list = prs.slides._sldIdLst
    for sld_id in list(id_list):
        prs.part.drop_rel(sld_id.rId)
        id_list.remove(sld_id)


def _pick_layout(prs, layout_name: Optional[str]):
    layouts = list(prs.slide_layouts)
    if layout_name:
        for layout in layouts:
            if layout.name.lower() == layout_name.lower():
                return layout
    for layout in layouts:
        if any(h in layout.name.lower() for h in _BLANK_HINTS):
            return layout
    return min(layouts, key=lambda l: len(l.placeholders))


def _target_box(prs, layout, margin: int) -> tuple[int, int, int, int]:
    for ph in layout.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
            return ph.left, ph.top, ph.width, ph.height
    return margin, margin, prs.slide_width - 2 * margin, prs.slide_height - 2 * margin


def _fit(img_w: int, img_h: int, box: tuple[int, int, int, int]) -> tuple[int, int, int, int]:
    left, top, width, height = box
    scale = min(width / img_w, height / img_h)
    w, h = int(img_w * scale), int(img_h * scale)
    return left + (width - w) // 2, top + (height - h) // 2, w, h


def build_pptx(
    images: list[Path],
    output: Path,
    template: Optional[Path] = None,
    notes: Optional[dict[int, str]] = None,
    layout_name: Optional[str] = None,
    margin_cm: float = 0.0,
) -> Path:
    """Crée une présentation avec une capture par diapo.

    ``notes`` : texte des notes du présentateur par numéro de diapo (1 = première image).
    Avec un modèle, la capture est placée dans l'espace réservé « image » de la
    disposition s'il existe, sinon centrée sur la diapo.
    """
    if not images:
        raise ValueError("Aucune diapo à exporter.")
    prs = open_template(template)
    _remove_existing_slides(prs)
    if template is None:
        with Image.open(images[0]) as first:
            w, h = first.size
        prs.slide_width = Emu(12192000)  # 33,867 cm (16:9 standard)
        prs.slide_height = Emu(int(12192000 * h / w))
    layout = _pick_layout(prs, layout_name)
    box = _target_box(prs, layout, int(Emu(int(margin_cm * 360000))))

    for number, image_path in enumerate(images, start=1):
        slide = prs.slides.add_slide(layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        with Image.open(image_path) as img:
            iw, ih = img.size
        left, top, width, height = _fit(iw, ih, box)
        slide.shapes.add_picture(str(image_path), left, top, width, height)
        if notes and notes.get(number):
            slide.notes_slide.notes_text_frame.text = notes[number]

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output


REPORT_LABELS = {
    "fr": {
        "title": "Compte-rendu de réunion",
        "before": "Avant la première diapo",
        "slide": "Diapo",
        "empty": "(rien de transcrit pendant cette diapo)",
        "summary": "Résumé",
        "summary_note": "Phrases clés extraites automatiquement de la transcription (sans reformulation).",
        "keywords": "Mots-clés",
        "brief": "En bref : ",
        "ai_note": "Rédigé par l'IA locale ({model}) à partir de la transcription — à vérifier.",
        "spoken": "Phrases clés réellement prononcées",
    },
    "en": {
        "title": "Meeting report",
        "before": "Before the first slide",
        "slide": "Slide",
        "empty": "(nothing transcribed during this slide)",
        "summary": "Summary",
        "summary_note": "Key sentences automatically extracted from the transcript (not rephrased).",
        "keywords": "Keywords",
        "brief": "In brief: ",
        "ai_note": "Written by the local AI ({model}) from the transcript — to be checked.",
        "spoken": "Key sentences actually spoken",
    },
}


def build_report(
    sections: list[SlideSection],
    slide_images: dict[int, Path],
    output: Path,
    title: Optional[str] = None,
    subtitle: str = "",
    offset: float = 0.0,
    summary_sentences: Optional[list[str]] = None,
    keywords: Optional[list[str]] = None,
    language: str = "fr",
    ai_summary: str = "",
    ai_model: str = "",
    slide_summaries: Optional[dict] = None,
) -> Path:
    """Document Word : chaque diapo suivie de ce qui a été dit pendant qu'elle était
    affichée, puis le résumé."""
    import docx
    from docx.shared import Cm, RGBColor

    labels = REPORT_LABELS.get(language, REPORT_LABELS["fr"])
    doc = docx.Document()
    doc.add_heading(title or labels["title"], level=0)
    if subtitle:
        doc.add_paragraph(subtitle)

    briefed: set = set()
    for section in sections:
        end = f" → {format_timestamp(section.end)}" if section.end is not None else ""
        span = f"{format_timestamp(section.start)}{end}"
        if section.slide_index is None:
            doc.add_heading(f"{labels['before']} ({span})", level=2)
        else:
            doc.add_heading(f"{labels['slide']} {section.slide_index} ({span})", level=2)
            image = slide_images.get(section.slide_index)
            if image and Path(image).exists():
                doc.add_picture(str(image), width=Cm(16))
            brief = (slide_summaries or {}).get(str(section.slide_index))
            if brief and section.slide_index not in briefed:
                briefed.add(section.slide_index)
                p = doc.add_paragraph()
                p.add_run(labels["brief"]).bold = True
                p.add_run(brief).italic = True
        if not section.segments:
            p = doc.add_paragraph(labels["empty"])
            p.runs[0].italic = True
        for seg in section.segments:
            p = doc.add_paragraph()
            stamp = p.add_run(f"[{format_timestamp(seg.start + offset)}] ")
            stamp.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            stamp.font.size = Pt(9)
            if seg.speaker:
                p.add_run(f"{seg.speaker} : ").bold = True
            p.add_run(seg.text)

    if ai_summary:
        doc.add_heading(labels["summary"], level=1)
        note = doc.add_paragraph(labels["ai_note"].format(model=ai_model))
        note.runs[0].italic = True
        _add_markdown(doc, ai_summary)
    if summary_sentences:
        doc.add_heading(labels["spoken"] if ai_summary else labels["summary"], level=1 if not ai_summary else 2)
        note = doc.add_paragraph(labels["summary_note"])
        note.runs[0].italic = True
        for sentence in summary_sentences:
            doc.add_paragraph(sentence, style="List Bullet")
        if keywords:
            p = doc.add_paragraph()
            p.add_run(f"{labels['keywords']} : ").bold = True
            p.add_run(", ".join(keywords))

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output


def _add_markdown(doc, text: str) -> None:
    """Rend un texte simple : « ## titre », « - puce », **gras**, paragraphes."""
    import re

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = re.match(r"^#{1,6}\s*(.+)$", line)
        bullet = re.match(r"^[-*•]\s+(.+)$", line)
        if heading:
            doc.add_heading(heading.group(1).strip("* "), level=2)
            continue
        p = doc.add_paragraph(style="List Bullet") if bullet else doc.add_paragraph()
        content = bullet.group(1) if bullet else line
        for k, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", content)):
            if chunk:
                p.add_run(chunk).bold = k % 2 == 1
