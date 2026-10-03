"""Lecture des transcriptions horodatées (export Plaud .txt / .srt / .vtt / .docx)
et association de chaque passage à la diapo affichée à ce moment-là."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

TS = r"(?:\d{1,2}:)?\d{1,2}:\d{2}(?:[.,]\d{1,3})?"
_ts_start = re.compile(rf"^\s*[\[(]?(?P<ts>{TS})[\])]?\s*(?:[-–:|]\s*)?(?P<rest>.*)$")
_ts_end = re.compile(rf"^\s*(?P<speaker>.*?\S)\s*[\[(]?(?P<ts>{TS})[\])]?\s*$")
_cue = re.compile(rf"(?P<start>{TS})\s*-->\s*(?P<end>{TS})")
_speaker_prefix = re.compile(r"^(?P<speaker>[^:]{1,40}?)\s*:\s+(?P<text>.+)$")


@dataclass
class Segment:
    start: float
    text: str
    speaker: str = ""
    end: Optional[float] = None

    def line(self, offset: float = 0.0) -> str:
        from .slides import format_timestamp

        who = f"{self.speaker} : " if self.speaker else ""
        return f"[{format_timestamp(self.start + offset)}] {who}{self.text}"


def parse_timestamp(value: str) -> float:
    value = value.replace(",", ".")
    parts = value.split(":")
    seconds = float(parts[-1])
    minutes = int(parts[-2]) if len(parts) >= 2 else 0
    hours = int(parts[-3]) if len(parts) >= 3 else 0
    return hours * 3600 + minutes * 60 + seconds


def read_text(path: Path) -> str:
    path = Path(path)
    if path.suffix.lower() == ".docx":
        import docx

        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_transcript(text: str) -> list[Segment]:
    if _cue.search(text):
        return _parse_cues(text)
    return _parse_generic(text)


def load_transcript(path: Path) -> list[Segment]:
    segments = parse_transcript(read_text(path))
    if not segments:
        raise ValueError(
            "Aucun horodatage trouvé dans la transcription. "
            "Exportez-la depuis Plaud avec les horodatages (format TXT ou SRT)."
        )
    return segments


def _split_speaker(text: str) -> tuple[str, str]:
    m = _speaker_prefix.match(text)
    if m and not re.search(r"\d{2}:\d{2}", m.group("speaker")):
        return m.group("speaker").strip(), m.group("text").strip()
    return "", text.strip()


def _parse_cues(text: str) -> list[Segment]:
    segments: list[Segment] = []
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        lines = [l.strip() for l in block.strip().split("\n") if l.strip()]
        for i, line in enumerate(lines):
            m = _cue.search(line)
            if not m:
                continue
            body = " ".join(lines[i + 1 :])
            body = re.sub(r"<[^>]+>", "", body).strip()
            if body:
                speaker, body = _split_speaker(body)
                segments.append(
                    Segment(parse_timestamp(m["start"]), body, speaker, parse_timestamp(m["end"]))
                )
            break
    return segments


def _parse_generic(text: str) -> list[Segment]:
    segments: list[Segment] = []
    current: Optional[Segment] = None
    buffer: list[str] = []

    def flush() -> None:
        if current is not None:
            body = " ".join(buffer).strip()
            if body:
                current.text = (current.text + " " + body).strip() if current.text else body
            if current.text:
                segments.append(current)

    for raw_line in text.replace("\r\n", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        start = _ts_start.match(line)
        end = None if start else _ts_end.match(line)
        if start:
            flush()
            speaker, body = _split_speaker(start["rest"]) if start["rest"] else ("", "")
            current = Segment(parse_timestamp(start["ts"]), body, speaker)
            buffer = []
        elif end and len(end["speaker"].split()) <= 5:
            # Style « Intervenant 1  00:01:23 » suivi du texte sur les lignes suivantes.
            flush()
            current = Segment(parse_timestamp(end["ts"]), "", end["speaker"].rstrip(" :-|"))
            buffer = []
        elif current is not None:
            buffer.append(line)
    flush()
    return segments


def offset_from_clock(session_start: datetime, plaud_start: datetime) -> float:
    """Décalage (s) à ajouter aux horodatages Plaud pour les ramener au temps de la vidéo."""
    return (plaud_start - session_start).total_seconds()


@dataclass
class SlideSection:
    """Passage de la réunion pendant lequel une diapo était affichée."""

    slide_index: Optional[int]
    """None = avant la première diapo."""
    start: float
    end: Optional[float]
    segments: list[Segment] = field(default_factory=list)


def assign_segments(
    timeline: list[tuple[float, int]],
    segments: list[Segment],
    offset: float = 0.0,
    duration: Optional[float] = None,
) -> list[SlideSection]:
    """Répartit les passages de la transcription entre les diapos, dans l'ordre chronologique."""
    events = sorted(timeline)
    sections: list[SlideSection] = []
    if not events or events[0][0] > 0:
        sections.append(SlideSection(None, 0.0, events[0][0] if events else duration))
    for i, (t, index) in enumerate(events):
        end = events[i + 1][0] if i + 1 < len(events) else duration
        sections.append(SlideSection(index, t, end))

    starts = [s.start for s in sections]
    for seg in sorted(segments, key=lambda s: s.start):
        t = seg.start + offset
        pos = 0
        for j, st in enumerate(starts):
            if st <= t:
                pos = j
            else:
                break
        sections[pos].segments.append(seg)
    return [s for s in sections if s.slide_index is not None or s.segments]


def notes_by_slide(sections: list[SlideSection], offset: float = 0.0) -> dict[int, str]:
    """Texte des notes du présentateur, par numéro de diapo."""
    from .slides import format_timestamp

    notes: dict[int, list[str]] = {}
    for section in sections:
        if section.slide_index is None:
            continue
        lines = notes.setdefault(section.slide_index, [])
        lines.append(f"— Affichée à {format_timestamp(section.start)} —")
        lines.extend(seg.line(offset) for seg in section.segments)
    return {k: "\n".join(v) for k, v in notes.items()}
